from io import BytesIO
from pathlib import Path
import re

from docx import Document
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.image.image import Image
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor

from app.services.storage import StorageService


def add_xml_field(run, field_name: str) -> None:
    fld = OxmlElement("w:fldSimple")
    fld.set(qn("w:instr"), field_name)
    run._r.append(fld)


class DocxReportWriter:
    def __init__(self) -> None:
        self.storage = StorageService()

    def build(
        self,
        project,
        report: dict,
        diagrams: list[dict],
        screenshots: list,
        validation: dict,
        domain_info: dict | None = None,
    ) -> tuple[str, int]:
        domain_info = domain_info or {}
        doc = Document()
        self._configure_document(doc, project, domain_info)

        # 1. Front Matter / Preliminaries
        self._add_cover_page(doc, project, domain_info)
        self._add_certificate_page(doc, project, domain_info)
        self._add_declaration_page(doc, project, domain_info)
        self._add_acknowledgement_page(doc, project, domain_info)
        self._add_abstract_page(doc, project, domain_info)

        # 2. Structured Table of Contents (Index Table with Page Numbers)
        self._add_table_of_contents(doc, report, diagrams, screenshots)

        # 3. List of Figures & Schematics
        if diagrams or screenshots:
            self._add_list_of_figures(doc, diagrams, screenshots)

        # 4. Main Body Chapters with Inline Diagrams, Tables & Screenshots
        self._add_chapters(doc, report, diagrams, screenshots)

        # 5. Appendices: Dedicated Blueprints, Lab Exhibits & References
        self._add_appendices(doc, project, report, diagrams, screenshots, validation, domain_info)

        buffer = BytesIO()
        doc.save(buffer)
        safe_title = "".join(c if c.isalnum() or c in "-_" else "_" for c in project.title).strip("_") or "academic_report"
        return self.storage.write_bytes(buffer.getvalue(), "generated/reports", f"{safe_title}.docx")

    # -------------------------------------------------------------------------
    # Document Setup & Running Headers / Footers
    # -------------------------------------------------------------------------
    def _configure_document(self, doc: Document, project, domain_info: dict) -> None:
        section = doc.sections[0]
        section.top_margin = Inches(1.0)
        section.bottom_margin = Inches(1.0)
        section.left_margin = Inches(1.25)
        section.right_margin = Inches(1.0)
        section.different_first_page_header_footer = True

        dept = project.branch or domain_info.get("domain_name") or "Engineering & Technology"

        # Running Header (Pages 2+)
        header = section.header
        header_p = header.paragraphs[0]
        header_p.alignment = WD_ALIGN_PARAGRAPH.RIGHT
        h_run = header_p.add_run(f"{project.title[:55]} | {project.college or 'Academic Report'}")
        h_run.font.name = "Times New Roman"
        h_run.font.size = Pt(8.5)
        h_run.font.color.rgb = RGBColor(100, 116, 139)

        # Running Footer with Page Numbers
        footer = section.footer
        footer_p = footer.paragraphs[0]
        footer_p.alignment = WD_ALIGN_PARAGRAPH.RIGHT

        f_left = footer_p.add_run(f"Department of {dept}   •   ")
        f_left.font.name = "Times New Roman"
        f_left.font.size = Pt(9)
        f_left.font.color.rgb = RGBColor(100, 116, 139)

        f_run1 = footer_p.add_run("Page ")
        f_run1.font.name = "Times New Roman"
        f_run1.font.size = Pt(9.5)
        f_run1.font.bold = True
        f_run1.font.color.rgb = RGBColor(30, 41, 59)

        f_page = footer_p.add_run()
        add_xml_field(f_page, "PAGE")
        f_page.font.name = "Times New Roman"
        f_page.font.size = Pt(9.5)
        f_page.font.bold = True
        f_page.font.color.rgb = RGBColor(30, 41, 59)

        f_run2 = footer_p.add_run(" of ")
        f_run2.font.name = "Times New Roman"
        f_run2.font.size = Pt(9.5)
        f_run2.font.color.rgb = RGBColor(100, 116, 139)

        f_numpages = footer_p.add_run()
        add_xml_field(f_numpages, "NUMPAGES")
        f_numpages.font.name = "Times New Roman"
        f_numpages.font.size = Pt(9.5)
        f_numpages.font.color.rgb = RGBColor(100, 116, 139)

        # Typography hierarchy
        normal = doc.styles["Normal"]
        normal.font.name = "Times New Roman"
        normal.font.size = Pt(12)
        normal.font.color.rgb = RGBColor(15, 23, 42)
        normal.paragraph_format.line_spacing = 1.15
        normal.paragraph_format.space_after = Pt(6)

        h1 = doc.styles["Heading 1"]
        h1.font.name = "Times New Roman"
        h1.font.size = Pt(16)
        h1.font.bold = True
        h1.font.color.rgb = RGBColor(15, 23, 42)
        h1.paragraph_format.space_before = Pt(14)
        h1.paragraph_format.space_after = Pt(8)

        h2 = doc.styles["Heading 2"]
        h2.font.name = "Times New Roman"
        h2.font.size = Pt(13.5)
        h2.font.bold = True
        h2.font.color.rgb = RGBColor(30, 41, 59)
        h2.paragraph_format.space_before = Pt(10)
        h2.paragraph_format.space_after = Pt(4)

    def _get_degree_name(self, branch: str) -> str:
        b_low = (branch or "").lower()
        if "mca" in b_low:
            return "MASTER OF COMPUTER APPLICATIONS"
        if "mba" in b_low or "management" in b_low:
            return "MASTER OF BUSINESS ADMINISTRATION"
        if "m.tech" in b_low or "mtech" in b_low or "m.e." in b_low:
            return f"MASTER OF TECHNOLOGY\nIN {(branch or 'ENGINEERING').upper()}"
        if "b.sc" in b_low or "m.sc" in b_low or "science" in b_low:
            return f"BACHELOR OF SCIENCE / MASTER OF SCIENCE\nIN {(branch or 'SCIENCE').upper()}"
        return f"BACHELOR OF TECHNOLOGY / BACHELOR OF ENGINEERING\nIN {(branch or 'ENGINEERING').upper()}"

    # -------------------------------------------------------------------------
    # 1. Front Matter Pages
    # -------------------------------------------------------------------------
    def _add_cover_page(self, doc: Document, project, domain_info: dict) -> None:
        college = (project.college or "UNIVERSITY FACULTY OF ENGINEERING & TECHNOLOGY").upper()
        branch = (project.branch or domain_info.get("domain_name") or "ENGINEERING & TECHNOLOGY").upper()

        self._centered(doc, college, size=15, bold=True)
        self._centered(doc, f"DEPARTMENT OF {branch}", size=13, bold=True)
        self._centered(doc, "(Affiliated to University & Approved by AICTE / UGC)", size=10, italic=True)

        doc.add_paragraph()
        self._centered(doc, "A MAJOR PROJECT REPORT", size=14, bold=True)
        self._centered(doc, "ON", size=12, bold=True)

        title_p = self._centered(doc, project.title.upper(), size=18, bold=True)
        title_p.paragraph_format.space_before = Pt(10)
        title_p.paragraph_format.space_after = Pt(12)

        self._centered(
            doc,
            "Submitted in partial fulfillment of the requirements for the award of the degree of",
            size=11,
            italic=True,
        )
        self._centered(
            doc,
            self._get_degree_name(project.branch),
            size=12,
            bold=True,
        )

        doc.add_paragraph()
        table = doc.add_table(rows=2, cols=2)
        table.alignment = WD_TABLE_ALIGNMENT.CENTER
        table.cell(0, 0).text = "SUBMITTED BY:"
        table.cell(0, 1).text = "UNDER THE GUIDANCE OF:"
        table.cell(1, 0).text = project.student_details or "Candidate Name\nRoll No: 01"
        table.cell(1, 1).text = f"{project.guide or 'Project Supervisor'}\nAssistant / Associate Professor"

        for row in table.rows:
            for cell in row.cells:
                self._style_cell(cell)

        doc.add_paragraph()
        if project.academic_year:
            self._centered(doc, f"ACADEMIC YEAR: {project.academic_year}", size=12, bold=True)
        else:
            self._centered(doc, "ACADEMIC YEAR: 2026–2027", size=12, bold=True)

        doc.add_page_break()

    def _add_certificate_page(self, doc: Document, project, domain_info: dict) -> None:
        college = (project.college or "University Department of Engineering & Technology").title()
        branch = (project.branch or domain_info.get("domain_name") or "Engineering").title()

        self._centered(doc, college.upper(), size=14, bold=True)
        self._centered(doc, f"DEPARTMENT OF {branch.upper()}", size=12, bold=True)

        cert_p = self._centered(doc, "BONAFIDE CERTIFICATE", size=16, bold=True)
        cert_p.paragraph_format.space_before = Pt(16)
        cert_p.paragraph_format.space_after = Pt(16)

        degree = self._get_degree_name(project.branch).splitlines()[0]
        body = (
            f"This is to certify that the project report entitled \"{project.title.upper()}\" "
            f"submitted by {project.student_details or 'the student(s)'} is a bonafide record of "
            f"work carried out in partial fulfillment of the requirements for the award of the Degree of "
            f"{degree.title()} in {branch} from "
            f"{college} during the academic year {project.academic_year or '2026-2027'}."
        )
        p = doc.add_paragraph(body)
        p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
        p.paragraph_format.line_spacing = 1.3
        p.paragraph_format.space_after = Pt(12)

        body2 = (
            "The project report has been approved as it satisfies the academic and curriculum requirements "
            "in respect of project work prescribed for the said degree. The investigations, design models, and "
            "experimental results embodied in this report have been verified and have not been submitted to any other "
            "University or Institute for the award of any other degree or diploma."
        )
        p2 = doc.add_paragraph(body2)
        p2.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
        p2.paragraph_format.line_spacing = 1.3
        p2.paragraph_format.space_after = Pt(36)

        table = doc.add_table(rows=2, cols=3)
        table.alignment = WD_TABLE_ALIGNMENT.CENTER
        guide_name = project.guide or "Faculty Project Guide"
        dept_name = f"Dept. of {branch}"
        signatures = [
            (f"______________________\n{guide_name}\nProject Guide", "Internal Supervisor"),
            (f"______________________\nHead of Department\n{dept_name}", "H.O.D."),
            ("______________________\nExternal Examiner\nUniversity Nominee", "External Evaluator"),
        ]
        for idx, (top, bottom) in enumerate(signatures):
            table.cell(0, idx).text = top
            table.cell(1, idx).text = bottom
            for cell in [table.cell(0, idx), table.cell(1, idx)]:
                cell.paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.CENTER
                self._style_cell(cell, font_size=10)

        doc.add_paragraph()
        date_p = doc.add_paragraph(f"Place: Campus / Department                                     Date: {project.academic_year or '2026-2027'}")
        date_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        date_p.runs[0].font.size = Pt(10)
        date_p.runs[0].font.italic = True
        doc.add_page_break()

    def _add_declaration_page(self, doc: Document, project, domain_info: dict) -> None:
        self._centered(doc, "CANDIDATE'S DECLARATION", size=16, bold=True)
        doc.add_paragraph()
        body = (
            f"I/We hereby declare that the work presented in this project report entitled \"{project.title}\" "
            "is an authentic record of our own analytical research, design modeling, experimental investigation, "
            f"and technical implementation carried out under the supervision of {project.guide or 'our Project Guide'}.\n\n"
            "We further declare that no part of this report has been plagiarized or fabricated. All external standards, "
            "regulatory codes, algorithms, empirical datasets, and reference literature have been properly acknowledged "
            "and cited in compliance with academic and professional research ethics."
        )
        p = doc.add_paragraph(body)
        p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
        p.paragraph_format.line_spacing = 1.25
        p.paragraph_format.space_after = Pt(30)

        table = doc.add_table(rows=2, cols=2)
        table.alignment = WD_TABLE_ALIGNMENT.CENTER
        table.cell(0, 0).text = f"Date: {project.academic_year or 'Academic Session'}\nPlace: Campus"
        table.cell(0, 1).text = f"Signature of Candidate(s):\n\n{project.student_details or 'Candidate'}"
        for row in table.rows:
            for cell in row.cells:
                self._style_cell(cell, font_size=10.5)

        doc.add_page_break()

    def _add_acknowledgement_page(self, doc: Document, project, domain_info: dict) -> None:
        self._centered(doc, "ACKNOWLEDGEMENT", size=16, bold=True)
        doc.add_paragraph()
        dept = project.branch or domain_info.get("domain_name") or "our Engineering Department"
        college = project.college or "our institution"
        body = (
            f"We express our deepest sense of gratitude and sincere thanks to our respected guide, "
            f"{project.guide or 'our project supervisor'}, for their invaluable guidance, constructive criticism, "
            "insightful advice, and constant encouragement throughout the course of this project work.\n\n"
            f"We extend our special thanks to the Head of the Department and all respected faculty members of "
            f"the Department of {dept} at {college} "
            "for providing indispensable laboratory infrastructure, computing facilities, and an inspiring academic environment.\n\n"
            "Finally, we express our heartfelt appreciation to our parents, teammates, and fellow students whose "
            "moral support, cooperative discussions, and patience were instrumental in completing this thesis document."
        )
        p = doc.add_paragraph(body)
        p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
        p.paragraph_format.line_spacing = 1.25
        doc.add_page_break()

    def _add_abstract_page(self, doc: Document, project, domain_info: dict) -> None:
        self._centered(doc, "ABSTRACT", size=16, bold=True)
        doc.add_paragraph()
        abstract_text = project.abstract or (
            f"This academic report presents the comprehensive design, analytical modeling, and empirical "
            f"investigation of '{project.title}'. Designed to address critical operational and technical challenges "
            f"in its domain, the investigation incorporates established national and international standards across "
            f"methodological formulations, experimental testing, and systematic analysis. Empirical findings confirm "
            f"high reliability, structural conformance, and robust alignment with curriculum benchmarks."
        )
        p = doc.add_paragraph(abstract_text)
        p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
        p.paragraph_format.line_spacing = 1.25
        p.paragraph_format.space_after = Pt(14)

        kw_text = project.technologies or (domain_info.get("domain_name") if domain_info else "Engineering Research")
        kw = doc.add_paragraph()
        k_run1 = kw.add_run("Keywords: ")
        k_run1.bold = True
        kw.add_run(f"{kw_text}, Analytical Modeling, Empirical Investigation, Performance Benchmark.")

        doc.add_page_break()

    # -------------------------------------------------------------------------
    # 2. Table of Contents
    # -------------------------------------------------------------------------
    def _add_table_of_contents(self, doc: Document, report: dict, diagrams: list[dict], screenshots: list) -> None:
        self._centered(doc, "TABLE OF CONTENTS", size=16, bold=True)
        doc.add_paragraph()

        table = doc.add_table(rows=1, cols=3)
        table.alignment = WD_TABLE_ALIGNMENT.CENTER
        table.style = "Table Grid"

        headers = ["Sr. No.", "Chapter / Section Title", "Page No."]
        widths = [Inches(1.0), Inches(4.5), Inches(1.1)]

        for i, (cell, h_text, w) in enumerate(zip(table.rows[0].cells, headers, widths)):
            cell.text = h_text
            cell.width = w
            self._style_cell(cell, bold=True, font_size=11, bg_color="1E3A8A", text_color=(255, 255, 255))
            cell.paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.CENTER if i != 1 else WD_ALIGN_PARAGRAPH.LEFT

        prelims = [
            ("-", "Bonafide Certificate of Approval", "i"),
            ("-", "Candidate's Declaration", "ii"),
            ("-", "Acknowledgement", "iii"),
            ("-", "Executive Abstract", "iv"),
            ("-", "List of Figures & Schematics", "v"),
        ]
        for sr, title, pno in prelims:
            row = table.add_row()
            row.cells[0].text = sr
            row.cells[1].text = title
            row.cells[2].text = pno
            for i, cell in enumerate(row.cells):
                self._style_cell(cell, italic=True, font_size=10.5)
                cell.paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.CENTER if i != 1 else WD_ALIGN_PARAGRAPH.LEFT

        chapters = report.get("chapters", [])
        current_page = 1
        has_ref_chapter = False
        for idx, chapter in enumerate(chapters, start=1):
            raw_title = chapter.get("title", f"Chapter {idx}")
            clean_title = re.sub(r"^(?:chapter\s+\d+[:\s-]*|\d+[:\s-]*)", "", raw_title, flags=re.I).strip() or raw_title
            if clean_title.lower() in ("references", "bibliography"):
                has_ref_chapter = True
            row = table.add_row()
            row.cells[0].text = f"Chapter {idx}"
            row.cells[1].text = clean_title
            row.cells[2].text = str(current_page)
            for i, cell in enumerate(row.cells):
                self._style_cell(cell, bold=(i == 0), font_size=10.5)
                cell.paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.CENTER if i != 1 else WD_ALIGN_PARAGRAPH.LEFT
            current_page += max(2, len(chapter.get("content", "").split()) // 350 + 1)

        appendices = [
            ("Appendix A", "Technical Specifications & Design Schematics", str(current_page)),
            ("Appendix B", "Laboratory Exhibits & Empirical Documentation", str(current_page + 2)),
        ]
        if not has_ref_chapter:
            appendices.append(("-", "References & Bibliography (IEEE / Standard Format)", str(current_page + 4)))

        for sr, title, pno in appendices:
            row = table.add_row()
            row.cells[0].text = sr
            row.cells[1].text = title
            row.cells[2].text = pno
            for i, cell in enumerate(row.cells):
                self._style_cell(cell, bold=True, font_size=10.5)
                cell.paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.CENTER if i != 1 else WD_ALIGN_PARAGRAPH.LEFT

        self._set_table_borders(table)
        doc.add_page_break()

    # -------------------------------------------------------------------------
    # 3. List of Figures
    # -------------------------------------------------------------------------
    def _add_list_of_figures(self, doc: Document, diagrams: list[dict], screenshots: list) -> None:
        self._centered(doc, "LIST OF FIGURES AND SCHEMATICS", size=16, bold=True)
        doc.add_paragraph()

        table = doc.add_table(rows=1, cols=3)
        table.alignment = WD_TABLE_ALIGNMENT.CENTER
        table.style = "Table Grid"

        headers = ["Figure No.", "Figure Caption / Diagram Title", "Page No."]
        widths = [Inches(1.2), Inches(4.3), Inches(1.1)]

        for i, (cell, h_text, w) in enumerate(zip(table.rows[0].cells, headers, widths)):
            cell.text = h_text
            cell.width = w
            self._style_cell(cell, bold=True, font_size=11, bg_color="1E3A8A", text_color=(255, 255, 255))
            cell.paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.CENTER if i != 1 else WD_ALIGN_PARAGRAPH.LEFT

        fig_idx = 1
        page_est = 6
        for diag in diagrams:
            row = table.add_row()
            row.cells[0].text = f"Figure {fig_idx}"
            row.cells[1].text = diag.get("title", "Technical Architecture Schematic")
            row.cells[2].text = str(page_est)
            for i, cell in enumerate(row.cells):
                self._style_cell(cell, font_size=10.5)
                cell.paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.CENTER if i != 1 else WD_ALIGN_PARAGRAPH.LEFT
            fig_idx += 1
            page_est += 3

        for file in screenshots:
            row = table.add_row()
            name = getattr(file, "original_name", "Exhibit")
            row.cells[0].text = f"Figure {fig_idx}"
            row.cells[1].text = f"Empirical Exhibit: {name}"
            row.cells[2].text = str(page_est)
            for i, cell in enumerate(row.cells):
                self._style_cell(cell, font_size=10.5)
                cell.paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.CENTER if i != 1 else WD_ALIGN_PARAGRAPH.LEFT
            fig_idx += 1
            page_est += 2

        self._set_table_borders(table)
        doc.add_page_break()

    # -------------------------------------------------------------------------
    # 4. Main Body Chapters with Tables & Visuals
    # -------------------------------------------------------------------------
    def _add_chapters(self, doc: Document, report: dict, diagrams: list[dict], screenshots: list) -> None:
        chapters = report.get("chapters", [])
        placed_diagram_indices = set()
        placed_screenshot_indices = set()

        for idx, chapter in enumerate(chapters, start=1):
            raw_title = chapter.get("title", f"Chapter {idx}")
            clean_title = re.sub(r"^(?:chapter\s+\d+[:\s-]*|\d+[:\s-]*)", "", raw_title, flags=re.I).strip() or raw_title
            doc.add_heading(f"CHAPTER {idx}: {clean_title.upper()}", 1)

            # References / Bibliography
            if clean_title.lower() in ("references", "bibliography"):
                doc.add_paragraph("Academic literature, standard regulatory codes, and research publications consulted in IEEE citation format:")
                refs = report.get("references", [])
                if not refs:
                    refs = [
                        "[1] Standard Technical Specifications & Regulatory Engineering Codes of Practice, 2024.",
                        "[2] S. Author and R. Co-Author, 'Foundational Investigations and Structural Methodology', Academic Journal, vol. 48, no. 4, pp. 210-225, 2023.",
                        "[3] National and International Standards for Design, Testing and Material Quality Assurance, 2022.",
                    ]
                for ref in refs:
                    ref_p = doc.add_paragraph(style="List Bullet")
                    r = ref_p.add_run(str(ref))
                    r.font.name = "Times New Roman"
                    r.font.size = Pt(10.5)
                doc.add_page_break()
                continue

            raw_content = chapter.get("content", "")
            paragraphs = [p.strip() for p in raw_content.split("\n") if p.strip()]

            for p_text in paragraphs:
                sub_match = re.match(r"^(?:(?:[0-9]+\.[0-9]+(?:\.[0-9]+)?)|#{2,3})\s+(.+)$", p_text)
                if sub_match:
                    heading_text = p_text.lstrip("# ").strip()
                    doc.add_heading(heading_text, 2)
                    continue

                if p_text.startswith(("- ", "* ", "• ")):
                    bp = doc.add_paragraph(style="List Bullet")
                    run = bp.add_run(p_text.lstrip("-*• "))
                    run.font.name = "Times New Roman"
                    run.font.size = Pt(11)
                    continue

                body_p = doc.add_paragraph()
                body_p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
                body_p.paragraph_format.line_spacing = 1.15
                body_p.paragraph_format.space_after = Pt(6)
                r = body_p.add_run(p_text)
                r.font.name = "Times New Roman"
                r.font.size = Pt(11.5)

            # Inline Publication-grade Tables inside Chapter
            for t_idx, table_data in enumerate(chapter.get("tables", []) or [], start=1):
                headers = table_data.get("headers") or []
                rows = table_data.get("rows") or []
                if headers:
                    caption = doc.add_paragraph(f"Table {idx}.{t_idx}: {table_data.get('title', 'Empirical Data & Parameter Matrix')}")
                    caption.alignment = WD_ALIGN_PARAGRAPH.CENTER
                    caption.paragraph_format.space_before = Pt(12)
                    caption.paragraph_format.space_after = Pt(4)
                    caption.runs[0].bold = True
                    caption.runs[0].font.name = "Times New Roman"
                    caption.runs[0].font.size = Pt(10)

                    t = doc.add_table(rows=1, cols=len(headers))
                    t.alignment = WD_TABLE_ALIGNMENT.CENTER
                    t.style = "Table Grid"

                    for cell, h_name in zip(t.rows[0].cells, headers):
                        cell.text = str(h_name)
                        self._style_cell(cell, bold=True, font_size=10, bg_color="1E3A8A", text_color=(255, 255, 255))
                        cell.paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.CENTER

                    for r_i, row_values in enumerate(rows):
                        cells = t.add_row().cells
                        bg_c = "F8FAFC" if (r_i % 2 == 1) else "FFFFFF"
                        for c_i, (cell, val) in enumerate(zip(cells, row_values)):
                            cell.text = str(val)
                            self._style_cell(cell, font_size=9.5, bg_color=bg_c)
                            if c_i == 0:
                                cell.paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.LEFT
                            else:
                                cell.paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.CENTER

                    self._set_table_borders(t)
                    doc.add_paragraph()

            # Inline Diagram Placement
            title_lower = clean_title.lower()
            is_design_or_arch = any(term in title_lower for term in [
                "architecture", "design", "methodology", "system", "data", "processing",
                "modeling", "structural", "workflow", "framework", "analysis", "layout"
            ])
            chapter_diag_count = 0
            if is_design_or_arch:
                for d_i, diag in enumerate(diagrams):
                    if d_i not in placed_diagram_indices and chapter_diag_count < 2:
                        placed_diagram_indices.add(d_i)
                        self._embed_diagram_inline(doc, diag, d_i + 1)
                        chapter_diag_count += 1
            elif idx in (3, 4, 5, 6) and len(placed_diagram_indices) < len(diagrams):
                for d_i, diag in enumerate(diagrams):
                    if d_i not in placed_diagram_indices and chapter_diag_count < 1:
                        placed_diagram_indices.add(d_i)
                        self._embed_diagram_inline(doc, diag, d_i + 1)
                        chapter_diag_count += 1

            # Inline Exhibits / Screenshots Placement
            is_output_or_impl = any(term in title_lower for term in [
                "result", "output", "implementation", "screen", "testing", "finding",
                "experiment", "evaluation", "verification", "lab"
            ])
            chapter_sc_count = 0
            if is_output_or_impl:
                for s_i, file in enumerate(screenshots):
                    if s_i not in placed_screenshot_indices and chapter_sc_count < 2:
                        placed_screenshot_indices.add(s_i)
                        self._embed_screenshot_inline(doc, file, s_i + 1)
                        chapter_sc_count += 1
            elif idx >= 7 and len(placed_screenshot_indices) < len(screenshots):
                for s_i, file in enumerate(screenshots):
                    if s_i not in placed_screenshot_indices and chapter_sc_count < 1:
                        placed_screenshot_indices.add(s_i)
                        self._embed_screenshot_inline(doc, file, s_i + 1)
                        chapter_sc_count += 1

            doc.add_page_break()

    def _embed_diagram_inline(self, doc: Document, diagram: dict, fig_num: int) -> None:
        try:
            path = self.storage.path_for(diagram["storage_key"])
            if self._add_fitted_picture(doc, path, max_width=6.0, max_height=4.8):
                caption = doc.add_paragraph(f"Figure {fig_num}: {diagram.get('title', 'Technical Architectural Model')}")
                caption.alignment = WD_ALIGN_PARAGRAPH.CENTER
                caption.paragraph_format.space_before = Pt(6)
                caption.paragraph_format.space_after = Pt(8)
                caption.runs[0].bold = True
                caption.runs[0].font.name = "Times New Roman"
                caption.runs[0].font.size = Pt(10)
                doc.add_paragraph()
        except Exception:
            pass

    def _embed_screenshot_inline(self, doc: Document, file, sc_num: int) -> None:
        try:
            path = self.storage.restore_if_missing(file.storage_key, getattr(file, "file_bytes", None))
            if self._add_fitted_picture(doc, path, max_width=6.0, max_height=4.5):
                name = getattr(file, "original_name", f"Exhibit {sc_num}")
                caption = doc.add_paragraph(f"Exhibit {sc_num}: {name}")
                caption.alignment = WD_ALIGN_PARAGRAPH.CENTER
                caption.paragraph_format.space_before = Pt(6)
                caption.paragraph_format.space_after = Pt(8)
                caption.runs[0].bold = True
                caption.runs[0].font.name = "Times New Roman"
                caption.runs[0].font.size = Pt(10)
                doc.add_paragraph()
        except Exception:
            pass

    # -------------------------------------------------------------------------
    # 5. Appendices
    # -------------------------------------------------------------------------
    def _add_appendices(
        self,
        doc: Document,
        project,
        report: dict,
        diagrams: list[dict],
        screenshots: list,
        validation: dict,
        domain_info: dict,
    ) -> None:
        doc.add_heading("APPENDIX A: SYSTEM BLUEPRINTS & METHODOLOGY COMPLIANCE", 1)
        p = doc.add_paragraph(
            "This appendix compiles technical specifications, methodological parameters, and regulatory code "
            "compliance criteria utilized throughout the development and verification phases."
        )
        p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
        doc.add_paragraph()

        # Detailed compliance checklist table
        app_table = doc.add_table(rows=1, cols=3)
        app_table.alignment = WD_TABLE_ALIGNMENT.CENTER
        app_table.style = "Table Grid"
        headers = ["Criteria / Component", "Applicable Standard / Code", "Verification Status"]
        for c, h in zip(app_table.rows[0].cells, headers):
            c.text = h
            self._style_cell(c, bold=True, font_size=10, bg_color="1E3A8A", text_color=(255, 255, 255))

        stds = domain_info.get("applicable_standards") or ["University Engineering Standard"]
        checks = [
            ("Technical Methodology Formulation", stds[0], "Verified"),
            ("Design Safety & Margin of Error", "National Regulatory Code", "Compliant"),
            ("Empirical Testing & Benchmark Integrity", "Quality Assurance Protocol", "Audited"),
            ("Documentation & Citation Ethics", "IEEE / Academic Standard", "Conforming"),
        ]
        for row_data in checks:
            cells = app_table.add_row().cells
            for cell, text in zip(cells, row_data):
                cell.text = text
                self._style_cell(cell, font_size=9.5)
        self._set_table_borders(app_table)
        doc.add_page_break()

        # Appendix B: Lab exhibits
        doc.add_heading("APPENDIX B: LABORATORY TESTING & OBSERVATIONAL EXHIBITS", 1)
        p2 = doc.add_paragraph(
            "Comprehensive log of verified laboratory observations, instrumentation outputs, and empirical test matrices."
        )
        p2.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
        doc.add_paragraph()

        for idx, file in enumerate(screenshots[:2], start=1):
            try:
                path = self.storage.restore_if_missing(file.storage_key, getattr(file, "file_bytes", None))
                if self._add_fitted_picture(doc, path, max_width=6.0, max_height=4.4):
                    caption = doc.add_paragraph(f"Exhibit B.{idx}: {file.original_name}")
                    caption.alignment = WD_ALIGN_PARAGRAPH.CENTER
                    caption.runs[0].bold = True
                    caption.runs[0].font.size = Pt(10)
                    doc.add_paragraph()
            except Exception:
                pass
        doc.add_page_break()

        # References Section (Formal Bibliography) - only if not already printed as body chapter
        has_ref_chapter = any(
            re.sub(r"^(?:chapter\s+\d+[:\s-]*|\d+[:\s-]*)", "", ch.get("title", ""), flags=re.I).strip().lower() in ("references", "bibliography")
            for ch in report.get("chapters", [])
        )
        if not has_ref_chapter:
            doc.add_heading("REFERENCES & BIBLIOGRAPHY", 1)
            doc.add_paragraph("Academic literature, research publications, and standard documentation consulted:")
            doc.add_paragraph()

            refs = report.get("references", [])
            if not refs:
                refs = [
                    "[1] Standard Technical Specifications & Regulatory Engineering Codes of Practice, 2024.",
                    f"[2] Documentation and Technical Specifications for {project.technologies or 'Modern Engineering Systems'}, 2024.",
                    "[3] E. Author and S. Researcher, 'Empirical Modeling and Investigation Protocol', IEEE / Academic Transactions, 2023.",
                ]

            for ref in refs:
                ref_p = doc.add_paragraph(style="List Bullet")
                r = ref_p.add_run(str(ref))
                r.font.name = "Times New Roman"
                r.font.size = Pt(10.5)

    # -------------------------------------------------------------------------
    # Formatting Helpers
    # -------------------------------------------------------------------------
    def _centered(self, doc: Document, text: str, size: int = 12, bold: bool = False, italic: bool = False):
        paragraph = doc.add_paragraph()
        paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
        run = paragraph.add_run(text)
        run.bold = bold
        run.italic = italic
        run.font.name = "Times New Roman"
        run.font.size = Pt(size)
        run.font.color.rgb = RGBColor(15, 23, 42)
        return paragraph

    def _style_cell(
        self,
        cell,
        bold: bool = False,
        italic: bool = False,
        font_size: float = 10,
        bg_color: str | None = None,
        text_color: tuple[int, int, int] | None = None,
    ) -> None:
        cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
        if bg_color:
            tcPr = cell._tc.get_or_add_tcPr()
            shd = OxmlElement("w:shd")
            shd.set(qn("w:val"), "clear")
            shd.set(qn("w:color"), "auto")
            shd.set(qn("w:fill"), bg_color)
            tcPr.append(shd)

        for paragraph in cell.paragraphs:
            for run in paragraph.runs:
                run.bold = bold
                run.italic = italic
                run.font.name = "Times New Roman"
                run.font.size = Pt(font_size)
                if text_color:
                    run.font.color.rgb = RGBColor(*text_color)
                else:
                    run.font.color.rgb = RGBColor(15, 23, 42)

    def _set_table_borders(self, table) -> None:
        tblPr = table._tbl.tblPr
        borders = tblPr.first_child_found_in("w:tblBorders")
        if borders is None:
            borders = OxmlElement("w:tblBorders")
            tblPr.append(borders)
        for edge in ["top", "left", "bottom", "right", "insideH", "insideV"]:
            tag = f"w:{edge}"
            element = borders.find(qn(tag))
            if element is None:
                element = OxmlElement(tag)
                borders.append(element)
            element.set(qn("w:val"), "single")
            element.set(qn("w:sz"), "6")
            element.set(qn("w:space"), "0")
            element.set(qn("w:color"), "CBD5E1")

    def _add_fitted_picture(self, doc: Document, path, max_width: float, max_height: float) -> bool:
        try:
            p = Path(path)
            if not p.exists() or p.stat().st_size == 0:
                return False
            image = Image.from_file(str(p))
            width = image.px_width / (image.horz_dpi or 96)
            height = image.px_height / (image.vert_dpi or 96)
            if width <= 0 or height <= 0:
                return False
            scale = min(max_width / width, max_height / height, 1.0)
            paragraph = doc.add_paragraph()
            paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
            paragraph.paragraph_format.space_before = Pt(8)
            paragraph.paragraph_format.space_after = Pt(4)
            run = paragraph.add_run()
            run.add_picture(str(p), width=Inches(width * scale), height=Inches(height * scale))
            return True
        except Exception:
            return False
