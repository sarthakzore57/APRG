import csv
import logging
import mimetypes
from pathlib import Path
import xml.etree.ElementTree as ET
import zipfile

import fitz
from docx import Document

logger = logging.getLogger(__name__)

TEXT_EXTENSIONS = {
    ".py", ".js", ".jsx", ".ts", ".tsx", ".java", ".cs", ".cpp", ".c", ".php",
    ".rb", ".go", ".rs", ".sql", ".html", ".css", ".json", ".md", ".txt",
    ".r", ".m", ".xml", ".yaml", ".yml", ".ini", ".sh", ".bat", ".tex", ".log",
}


def extract_csv(path: Path) -> str:
    try:
        with path.open("r", encoding="utf-8", errors="replace") as f:
            reader = csv.reader(f)
            rows = [
                " | ".join(cell.strip() for cell in row)
                for idx, row in enumerate(reader)
                if idx < 120 and any(row)
            ]
            return f"Tabular Dataset ({path.name}):\n" + "\n".join(rows)
    except Exception as exc:
        return f"CSV dataset ({path.name}): {exc}"


def extract_xlsx(path: Path) -> str:
    try:
        with zipfile.ZipFile(path, "r") as z:
            shared_strings: list[str] = []
            if "xl/sharedStrings.xml" in z.namelist():
                tree = ET.fromstring(z.read("xl/sharedStrings.xml"))
                for elem in tree.iter():
                    if elem.tag.endswith("}t") and elem.text:
                        shared_strings.append(elem.text)

            sheet_names = [n for n in z.namelist() if n.startswith("xl/worksheets/sheet") and n.endswith(".xml")]
            out: list[str] = []
            for sheet_name in sheet_names[:3]:
                tree = ET.fromstring(z.read(sheet_name))
                rows: list[str] = []
                for row_el in tree.iter():
                    if row_el.tag.endswith("}row"):
                        row_vals: list[str] = []
                        for cell_el in row_el:
                            if cell_el.tag.endswith("}c"):
                                cell_type = cell_el.get("t")
                                v_el = None
                                for child in cell_el:
                                    if child.tag.endswith("}v"):
                                        v_el = child
                                        break
                                val = v_el.text if v_el is not None and v_el.text else ""
                                if cell_type == "s" and val.isdigit() and int(val) < len(shared_strings):
                                    val = shared_strings[int(val)]
                                row_vals.append(val)
                        if any(row_vals):
                            rows.append(" | ".join(row_vals))
                if rows:
                    out.append(f"Worksheet ({sheet_name}):\n" + "\n".join(rows[:100]))
            return "\n\n".join(out) if out else f"Excel spreadsheet: {path.name}"
    except Exception as exc:
        return f"Excel spreadsheet ({path.name}): could not parse XML ({exc})"


def extract_image(path: Path) -> str:
    suffix = path.suffix.lower()
    mime = "image/png"
    if suffix in [".jpg", ".jpeg"]:
        mime = "image/jpeg"
    elif suffix == ".webp":
        mime = "image/webp"

    try:
        from app.services.ai_provider import get_ai_provider
        ai = get_ai_provider()
        data = path.read_bytes()
        if 0 < len(data) <= 20 * 1024 * 1024:
            description = ai.describe_image(data, mime_type=mime)
            if description and "Image exhibit provided as project empirical documentation." not in description:
                return (
                    f"Visual Project Artifact ({path.name}):\n"
                    f"Computer Vision Inspection & Empirical Findings:\n{description.strip()}"
                )
    except Exception as exc:
        logger.debug("Image AI analysis skipped for %s: %s", path.name, exc)

    return f"Visual Project Artifact ({path.name}): Student-provided figure, drawing, or photograph."


def extract_text(path: Path) -> str:
    suffix = path.suffix.lower()
    if suffix == ".pdf":
        with fitz.open(path) as pdf:
            return "\n".join(page.get_text("text") for page in pdf)
    if suffix == ".docx":
        doc = Document(path)
        paragraphs_text = "\n".join(p.text for p in doc.paragraphs if p.text.strip())
        tables_text = []
        for table in doc.tables:
            for row in table.rows:
                tables_text.append(" | ".join(c.text.strip() for c in row.cells if c.text.strip()))
        return paragraphs_text + ("\n\nDocx Tables:\n" + "\n".join(tables_text) if tables_text else "")
    if suffix == ".csv":
        return extract_csv(path)
    if suffix in [".xlsx", ".xls"]:
        return extract_xlsx(path)
    if suffix in TEXT_EXTENSIONS:
        return path.read_text(encoding="utf-8", errors="replace")
    if suffix in {".png", ".jpg", ".jpeg", ".webp"}:
        return extract_image(path)
    return f"Unsupported text extraction for {mimetypes.guess_type(path.name)[0] or suffix}."


def clean_text(text: str) -> str:
    lines = [" ".join(line.split()) for line in text.replace("\x00", " ").splitlines()]
    return "\n".join(line for line in lines if line).strip()


def chunk_text(text: str, chunk_size: int = 1200, overlap: int = 160) -> list[str]:
    cleaned = clean_text(text)
    chunks: list[str] = []
    start = 0
    while start < len(cleaned):
        end = min(start + chunk_size, len(cleaned))
        chunk = cleaned[start:end].strip()
        if chunk:
            chunks.append(chunk)
        if end == len(cleaned):
            break
        start = max(0, end - overlap)
    return chunks
